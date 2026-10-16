"""
Classify literature-review articles by Country, Context, Vehicle, Model and Theme.

Classifies the same text as the "Combined text" column of the Results sheet in
Literature review.xlsx (title, abstract, author keywords, index keywords), but rebuilt
from those columns with publisher copyright/licence statements removed from the
abstract, and tags each article with every category whose terms appear in the text.

Classification sheets (Country, Context, Vehicle, Model, Theme):
  * each row is a category; the first cell is the category name (and is itself a term)
  * every non-empty cell in the row is an OR term
  * a cell holding a comma-separated list contributes each item as an OR term
    (commas inside brackets, e.g. "GM(1,1)", do not split)

Matching rules:
  * case-insensitive, whole-word (a term never matches inside a longer word)
  * simple plurals are accepted ("bus" -> "buses", "simulation" -> "simulations")
  * spaces and hyphens inside a term are interchangeable ("agent based" == "agent-based")
  * fully capitalised acronyms (e.g. "AI", "TOPSIS") are matched case-sensitively;
    every other term is case-insensitive

Outputs:
  * new sheet "Classification" (one row per article, 0/1 indicator per category,
    plus a readable list of matched categories and matched terms per dimension, and a
    "Mentions model*" 0/1 flag for any word starting with "model" in title/abstract/keywords)
  * new sheet "Classification correlations" (phi coefficient, Fisher exact p,
    Benjamini-Hochberg q for every category pair)
  * PNG charts in ./classification_outputs/
  * "EV Literature Explorer.html": interactive cross-filtering page built from
    explorer_template.html (kept next to this script) with this run's classifications embedded

Usage:
  python classify_literature.py [--input "Literature review.xlsx"] [--output OUT.xlsx]
                                [--alpha 0.05] [--top 30]
                                [--explorer PAGE.html] [--explorer-template TEMPLATE.html] [--no-explorer]
"""

import argparse
import json
import re
import shutil
import sys
import tempfile
from itertools import combinations
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import openpyxl
import pandas as pd
from matplotlib.colors import LinearSegmentedColormap
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from scipy.stats import fisher_exact

DIMENSIONS = ["Country", "Context", "Vehicle", "Model", "Theme"]
TEXT_COLUMN = "Combined text"
TEXT_PARTS = ["Title", "Abstract", "Author Keywords", "Index Keywords"]  # what "Combined text" concatenates
ID_COLUMNS = ["Authors", "Title", "Year", "Source title", "DOI", "EID"]
MODEL_FLAG = "Mentions model*"  # 0/1 column in the Classification sheet
CLASS_SHEET = "Classification"
CORR_SHEET = "Classification correlations"

# Chart palette (validated reference palette: categorical slots 1-4, blue sequential,
# blue<->red diverging with a gray midpoint).
DIM_COLOURS = {"Country": "#2a78d6", "Context": "#eb6834", "Vehicle": "#1baf7a", "Model": "#eda100",
               "Theme": "#e87ba4"}
NEUTRAL = "#b5b3ad"
INK, INK_2, INK_3 = "#0b0b0b", "#52514e", "#8a8984"
SURFACE, GRID = "#fcfcfb", "#e6e5e0"
SEQ = LinearSegmentedColormap.from_list(
    "seq_blue", ["#f4f8fd", "#cde2fb", "#9ec5f4", "#6da7ec", "#3987e5", "#256abf", "#184f95", "#0d366b"])
DIV = LinearSegmentedColormap.from_list(
    "div_blue_red", ["#184f95", "#3987e5", "#9ec5f4", "#f0efec", "#f3a9a8", "#e34948", "#a32727"])

plt.rcParams.update({
    "font.family": "sans-serif", "font.size": 9,
    "axes.edgecolor": GRID, "axes.labelcolor": INK_2, "axes.titlecolor": INK,
    "axes.titlesize": 11, "axes.titleweight": "bold", "axes.titlelocation": "left",
    "xtick.color": INK_2, "ytick.color": INK_2, "text.color": INK,
    "figure.facecolor": SURFACE, "axes.facecolor": SURFACE, "savefig.facecolor": SURFACE,
    "axes.spines.top": False, "axes.spines.right": False,
})


# --------------------------------------------------------------------------- terms
QUOTES = str.maketrans({"“": '"', "”": '"', "‘": "'", "’": "'", "‐": "-",
                        "‑": "-", "–": "-", "—": "-"})
SCOPUS_SYNTAX = re.compile(r"\)|\b(?:AND|OR|NOT)\b|TITLE-ABS-KEY")


def clean_term(raw):
    """Strip quotes, brackets and any trailing Scopus query syntax from one term."""
    t = str(raw).translate(QUOTES)
    if SCOPUS_SYNTAX.search(t):
        t = SCOPUS_SYNTAX.split(t)[0]
    return t.strip().strip('"\'()').strip()


def cell_terms(value):
    if value is None or (isinstance(value, float) and np.isnan(value)):
        return []
    # split on commas, except those inside brackets such as "GM(1,1)"
    terms = [clean_term(part) for part in re.split(r",(?![^()]*\))", str(value))]
    return [t for t in terms if t]


def term_pattern(term):
    """Whole-word regex for a term with optional plural and flexible space/hyphen."""
    words = re.split(r"[\s\-]+", term)
    body = r"[\s\-]*".join(re.escape(w) for w in words)
    plural = r"(?:e?s)?"
    if len(words[-1]) > 2 and words[-1][-1].lower() == "y" and words[-1][-2].lower() not in "aeiou":
        body, plural = body[:-1], r"(?:y|ies)"  # lorry -> lorries
    flags = 0 if term.isupper() else re.IGNORECASE  # fully capitalised acronyms are case-sensitive
    return re.compile(rf"(?<![^\W_]){body}{plural}(?![^\W_])", flags)


def load_categories(xlsx, sheet):
    """Return {category: [terms]} for a classification sheet (rows = categories)."""
    raw = pd.read_excel(xlsx, sheet_name=sheet, header=None, dtype=object)
    categories = {}
    for _, row in raw.iterrows():
        cells = [c for c in row.tolist() if not (c is None or (isinstance(c, float) and np.isnan(c)))]
        if not cells:
            continue
        name_terms = cell_terms(cells[0])
        if not name_terms:
            continue
        name = name_terms[0]
        terms = []
        for c in cells:
            terms.extend(cell_terms(c))
        seen, unique = set(), []
        for t in terms:
            if t.lower() not in seen:
                seen.add(t.lower())
                unique.append(t)
        categories[name] = unique
    return categories


# --------------------------------------------------------------------------- classification
def classify(texts, categories):
    """Return (indicator DataFrame, matched-terms Series) for one dimension."""
    patterns = {cat: [(t, term_pattern(t)) for t in terms] for cat, terms in categories.items()}
    flags = {cat: [] for cat in categories}
    matched = []
    for text in texts:
        text = str(text).translate(QUOTES) if isinstance(text, str) else ""
        hits = []
        for cat, pats in patterns.items():
            found = [t for t, p in pats if p.search(text)]
            flags[cat].append(int(bool(found)))
            if found:
                hits.append(f"{cat}: {', '.join(found)}")
        matched.append("; ".join(hits))
    return pd.DataFrame(flags, index=texts.index), pd.Series(matched, index=texts.index)


# A copyright statement starts with "©" followed by a year, capital or bracket ("© 2024 Elsevier",
# "©The Author(s)"), or with "Copyright" followed by (c)/©/a year/"by"/"for this". A lone "©" before
# lower-case text is a stray character inside the abstract and is left alone.
COPYRIGHT_START = re.compile(
    r"©\s*(?=[\d(A-Z])|\bCopyright\b(?=\s*(?:\(c\)|©|\d{4}|by\b|for this\b|[A-Z]))"
    # licence/publisher sentences without a © sign (case-sensitive, so ordinary prose is not caught)
    r"|\b(?:Published (?:by|on behalf|under)|Publishing services by|This (?:is an [Oo]pen[- ][Aa]ccess|"
    r"work is licensed|article is licensed|article is distributed|article published as an Open Access))\b"
    r"|\bAll rights(?:,[^.;]{0,80},)? (?:are )?reserved\b")
# Further sentences that belong to the statement (rights, licence, publisher lines)
COPYRIGHT_TAIL = re.compile(r"rights|licen[cs]|creative commons|open access|permits|reproduction|https?://"
                            r"|copyright|©|published by|terms and conditions|distributed under", re.I)
SENTENCE_END = re.compile(r"[.;](?=\s|$)")


def strip_copyright(text):
    """Remove publisher copyright/licence statements, which otherwise produce false matches
    (e.g. 'text and data mining, AI training' -> AI; 'Springer Nature India' -> India)."""
    for _ in range(5):  # a record can carry more than one statement
        m = COPYRIGHT_START.search(text)
        if not m:
            break
        start = m.start()
        if not text[start] == "©" and not text.startswith("Copyright", start):
            # a licence/rights sentence: remove it from the start of its sentence
            prev = [s.end() for s in SENTENCE_END.finditer(text, 0, start)]
            start = prev[-1] if prev else 0
        end = m.end()
        first = True
        while end < len(text):
            stop = SENTENCE_END.search(text, end)
            stop = stop.end() if stop else len(text)
            if not first and not COPYRIGHT_TAIL.search(text[end:stop]):
                break
            end, first = stop, False
        text = text[:start] + " " + text[end:]
    return text


def get_texts(results):
    """Title + abstract + author keywords + index keywords (the same concatenation as the
    'Combined text' formula), with copyright statements removed from the abstract."""
    if all(c in results for c in TEXT_PARTS):
        parts = results[TEXT_PARTS].fillna("").astype(str)
        parts["Abstract"] = parts["Abstract"].map(strip_copyright)
        return parts.agg(" ".join, axis=1)
    # Fallback: the source columns are missing, so use the Combined text column as it stands.
    print(f"Columns {TEXT_PARTS} not all present - using '{TEXT_COLUMN}' without copyright removal")
    return results[TEXT_COLUMN].fillna("").astype(str)


# --------------------------------------------------------------------------- statistics
def bh_fdr(p):
    p = np.asarray(p, dtype=float)
    q = np.full_like(p, np.nan)
    ok = ~np.isnan(p)
    pv = p[ok]
    order = np.argsort(pv)
    ranked = pv[order] * len(pv) / (np.arange(len(pv)) + 1)
    ranked = np.minimum.accumulate(ranked[::-1])[::-1]
    out = np.empty_like(pv)
    out[order] = np.clip(ranked, 0, 1)
    q[ok] = out
    return q


def pairwise_correlations(ind, dim_of):
    rows = []
    n = len(ind)
    for a, b in combinations(ind.columns, 2):
        x, y = ind[a].values, ind[b].values
        n11 = int((x & y).sum())
        n10 = int((x & ~y.astype(bool)).sum())
        n01 = int((~x.astype(bool) & y).sum())
        n00 = n - n11 - n10 - n01
        na, nb = n11 + n10, n11 + n01
        denom = np.sqrt(na * (n - na) * nb * (n - nb))
        phi = (n11 * n00 - n10 * n01) / denom if denom else np.nan
        p = fisher_exact([[n11, n10], [n01, n00]])[1] if denom else np.nan
        expected = na * nb / n
        rows.append({"Category A": a, "Dimension A": dim_of[a], "Category B": b, "Dimension B": dim_of[b],
                     "Articles A": na, "Articles B": nb, "Both": n11, "Expected both": round(expected, 1),
                     "Lift": round(n11 / expected, 2) if expected else np.nan,
                     "Phi": phi, "p (Fisher exact)": p})
    df = pd.DataFrame(rows)
    df["q (BH-FDR)"] = bh_fdr(df["p (Fisher exact)"])
    df["Cross-dimension"] = df["Dimension A"] != df["Dimension B"]
    return df


# --------------------------------------------------------------------------- charts
def short(cat):
    return cat.split(":", 1)[1].strip()


def chart_summary(ind, cats_by_dim, n, out):
    nrows = int(np.ceil(len(DIMENSIONS) / 2))
    tallest = max(len(cats_by_dim[d]) for d in DIMENSIONS)
    fig, axes = plt.subplots(nrows, 2, figsize=(14, nrows * max(4.8, 0.3 * tallest + 1.5)))
    for ax in axes.flat[len(DIMENSIONS):]:
        ax.set_visible(False)
    for ax, dim in zip(axes.flat, DIMENSIONS):
        cols = [f"{dim}: {c}" for c in cats_by_dim[dim]]
        counts = ind[cols].sum().rename(index=short).sort_values()
        none = int((ind[cols].sum(axis=1) == 0).sum())
        labels = ["(none)"] + counts.index.tolist()
        values = [none] + counts.values.tolist()
        colours = [NEUTRAL] + [DIM_COLOURS[dim]] * len(counts)
        y = np.arange(len(labels))
        ax.barh(y, values, color=colours, height=0.72, edgecolor=SURFACE, linewidth=2)
        ax.set_yticks(y, labels)
        ax.tick_params(axis="y", length=0)
        ax.xaxis.grid(True, color=GRID, linewidth=0.8)
        ax.set_axisbelow(True)
        ax.spines["left"].set_visible(False)
        vmax = max(values) or 1
        for yi, v in zip(y, values):
            ax.text(v + vmax * 0.01, yi, f"{v:,}  ({v / n:.1%})", va="center", fontsize=8, color=INK_2)
        ax.set_xlim(0, vmax * 1.28)
        multi = int((ind[cols].sum(axis=1) > 1).sum())
        ax.set_title(f"{dim}", pad=18)
        ax.text(0, 1.015,f"{n - none:,} of {n:,} articles classified; {multi:,} in more than one category",
                transform=ax.transAxes, fontsize=8, color=INK_3, va="bottom")
        ax.set_xlabel("Articles")
    fig.suptitle("Articles per classification category", x=0.01, ha="left", fontsize=14, fontweight="bold")
    fig.tight_layout(rect=(0, 0, 1, 0.97))
    fig.savefig(out, dpi=200)
    plt.close(fig)


def chart_by_year(ind, years, cats_by_dim, out):
    cols = [f"{d}: {c}" for d in DIMENSIONS for c in cats_by_dim[d]]
    table = ind[cols].groupby(years.values).sum().T
    table = table.loc[:, sorted(table.columns)]
    fig, ax = plt.subplots(figsize=(max(8, 0.42 * table.shape[1] + 4), 0.3 * len(cols) + 2))
    im = ax.imshow(table.values, aspect="auto", cmap=SEQ, interpolation="nearest")
    ax.set_xticks(range(table.shape[1]), [str(c) for c in table.columns], rotation=90)
    ax.set_yticks(range(len(cols)), cols)
    ax.tick_params(length=0)
    for s in ax.spines.values():
        s.set_visible(False)
    vmax = table.values.max() or 1
    for i in range(table.shape[0]):
        for j in range(table.shape[1]):
            v = table.values[i, j]
            if v:
                ax.text(j, i, f"{v}", ha="center", va="center", fontsize=6.5,
                        color="white" if v > vmax * 0.55 else INK)
    _dimension_rules(ax, cats_by_dim, axis="y")
    cb = fig.colorbar(im, ax=ax, fraction=0.025, pad=0.01)
    cb.set_label("Articles")
    cb.outline.set_visible(False)
    ax.set_title("Articles per category by publication year")
    fig.tight_layout()
    fig.savefig(out, dpi=200)
    plt.close(fig)


def chart_crosstabs(ind, cats_by_dim, out):
    pairs = list(combinations(DIMENSIONS, 2))
    nrows = int(np.ceil(len(pairs) / 3))
    fig, axes = plt.subplots(nrows, 3, figsize=(21, 7.5 * nrows))
    for ax in axes.flat[len(pairs):]:
        ax.set_visible(False)
    for ax, (d1, d2) in zip(axes.flat, pairs):
        c1 = [f"{d1}: {c}" for c in cats_by_dim[d1]]
        c2 = [f"{d2}: {c}" for c in cats_by_dim[d2]]
        m = ind[c1].T.values @ ind[c2].values
        im = ax.imshow(m, cmap=SEQ, aspect="auto")
        ax.set_xticks(range(len(c2)), [short(c) for c in c2], rotation=45, ha="right")
        ax.set_yticks(range(len(c1)), [short(c) for c in c1])
        ax.tick_params(length=0)
        for s in ax.spines.values():
            s.set_visible(False)
        vmax = m.max() or 1
        for i in range(m.shape[0]):
            for j in range(m.shape[1]):
                ax.text(j, i, f"{m[i, j]}", ha="center", va="center", fontsize=7,
                        color="white" if m[i, j] > vmax * 0.55 else INK)
        ax.set_title(f"{d1} × {d2}")
        ax.set_ylabel(d1)
        ax.set_xlabel(d2)
    fig.suptitle("Co-occurrence of categories across dimensions (article counts)",
                 x=0.01, ha="left", fontsize=14, fontweight="bold")
    fig.tight_layout(rect=(0, 0, 1, 0.97))
    fig.savefig(out, dpi=200)
    plt.close(fig)


def _dimension_rules(ax, cats_by_dim, axis="both"):
    edges = np.cumsum([len(cats_by_dim[d]) for d in DIMENSIONS])[:-1] - 0.5
    for e in edges:
        if axis in ("both", "y"):
            ax.axhline(e, color=INK_2, linewidth=1)
        if axis in ("both", "x"):
            ax.axvline(e, color=INK_2, linewidth=1)


def chart_correlation_matrix(ind, corr, cats_by_dim, alpha, out):
    cols = [f"{d}: {c}" for d in DIMENSIONS for c in cats_by_dim[d]]
    k = len(cols)
    idx = {c: i for i, c in enumerate(cols)}
    phi = np.full((k, k), np.nan)
    q = np.full((k, k), np.nan)
    for r in corr.itertuples(index=False):
        i, j = idx[r[0]], idx[r[2]]
        phi[i, j] = phi[j, i] = r.Phi
        q[i, j] = q[j, i] = r[-2]
    lower = np.tril(np.ones((k, k), dtype=bool), -1)
    sig = lower & (q < alpha)
    lim = np.nanmax(np.abs(phi[sig])) if sig.any() else 0.3
    lim = max(0.1, float(np.ceil(lim * 10) / 10))

    fig, ax = plt.subplots(figsize=(0.42 * k + 5, 0.42 * k + 3))
    faded = np.where(lower & ~sig, phi, np.nan)
    strong = np.where(sig, phi, np.nan)
    ax.imshow(faded, cmap=DIV, vmin=-lim, vmax=lim, alpha=0.18, interpolation="nearest")
    im = ax.imshow(strong, cmap=DIV, vmin=-lim, vmax=lim, interpolation="nearest")
    for i, j in zip(*np.where(sig)):
        v = phi[i, j]
        stars = "***" if q[i, j] < 0.001 else "**" if q[i, j] < 0.01 else "*"
        ax.text(j, i, f"{v:.2f}\n{stars}", ha="center", va="center", fontsize=5.5, linespacing=0.9,
                color="white" if abs(v) > lim * 0.6 else INK)
        ax.add_patch(plt.Rectangle((j - 0.5, i - 0.5), 1, 1, fill=False, edgecolor=INK, linewidth=0.6))
    ax.set_xticks(range(k), cols, rotation=90)
    ax.set_yticks(range(k), cols)
    ax.tick_params(length=0)
    for s in ax.spines.values():
        s.set_visible(False)
    _dimension_rules(ax, cats_by_dim)
    cb = fig.colorbar(im, ax=ax, fraction=0.03, pad=0.01)
    cb.set_label("Phi coefficient (− less / + more often together than chance)")
    cb.outline.set_visible(False)
    ax.set_title(f"Category correlations - significant pairs highlighted "
                 f"(BH-FDR q < {alpha}; * q<{alpha}, ** q<0.01, *** q<0.001; faded = not significant)",
                 fontsize=10)
    fig.tight_layout()
    fig.savefig(out, dpi=200 if k <= 40 else 130)  # keep very large matrices to a manageable file size
    plt.close(fig)


def chart_top_pairs(corr, alpha, top, out, cross_only=True):
    d = corr[(corr["q (BH-FDR)"] < alpha)]
    if cross_only:
        d = d[d["Cross-dimension"]]
    d = d.reindex(d["Phi"].abs().sort_values(ascending=False).index).head(top)
    if d.empty:
        print("No significant pairs to plot for", out.name)
        return
    d = d.iloc[::-1]
    labels = [f"{a}  ×  {b}" for a, b in zip(d["Category A"], d["Category B"])]
    colours = ["#e34948" if v > 0 else "#2a78d6" for v in d["Phi"]]
    fig, ax = plt.subplots(figsize=(11, 0.32 * len(d) + 1.8))
    y = np.arange(len(d))
    ax.barh(y, d["Phi"], color=colours, height=0.7, edgecolor=SURFACE, linewidth=2)
    ax.axvline(0, color=INK_2, linewidth=0.8)
    ax.set_yticks(y, labels)
    ax.tick_params(axis="y", length=0)
    ax.spines["left"].set_visible(False)
    ax.xaxis.grid(True, color=GRID, linewidth=0.8)
    ax.set_axisbelow(True)
    span = d["Phi"].abs().max()
    for yi, (v, both, lift, qv) in enumerate(zip(d["Phi"], d["Both"], d["Lift"], d["q (BH-FDR)"])):
        ax.text(v + (span * 0.02 if v >= 0 else -span * 0.02), yi,
                f"φ={v:.2f}  n={both}  lift={lift:.1f}  q={qv:.1g}",
                va="center", ha="left" if v >= 0 else "right", fontsize=7, color=INK_2)
    lo = min(0, d["Phi"].min()) - span * 0.95 * (d["Phi"].min() < 0)
    ax.set_xlim(lo, max(0, d["Phi"].max()) + span * 0.6)
    ax.set_xlabel("Phi coefficient   (red = co-occur more than chance, blue = less)")
    scope = "cross-dimension" if cross_only else "all"
    ax.set_title(f"Strongest significant {scope} correlations (top {len(d)}, BH-FDR q < {alpha})")
    fig.tight_layout()
    fig.savefig(out, dpi=200)
    plt.close(fig)


# --------------------------------------------------------------------------- workbook output
def write_sheet(wb, name, df, widths=None, freeze="A2"):
    if name in wb.sheetnames:
        del wb[name]
    ws = wb.create_sheet(name)
    ws.append(list(df.columns))
    for row in df.itertuples(index=False):
        ws.append([None if (isinstance(v, float) and np.isnan(v)) else
                   (v.item() if isinstance(v, np.generic) else v) for v in row])
    head_fill = PatternFill("solid", fgColor="DCE8F7")
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = head_fill
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for i, col in enumerate(df.columns, 1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(col, 12)
    ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions
    return ws


def save_workbook(src, dst, classification, corr, alpha):
    print(f"Writing sheets to {dst} ...")
    wb = openpyxl.load_workbook(src)
    widths = {"Authors": 25, "Title": 60, "Source title": 25, "DOI": 25, "EID": 20}
    widths.update({f"{d} categories": 22 for d in DIMENSIONS})
    widths.update({f"{d} matched terms": 35 for d in DIMENSIONS})
    ws = write_sheet(wb, CLASS_SHEET, classification, widths, freeze="C2")
    hit = PatternFill("solid", fgColor="CDE2FB")
    first_flag = list(classification.columns).index(next(c for c in classification.columns if ": " in c)) + 1
    last_flag = first_flag + sum(": " in c for c in classification.columns) - 1
    for row in ws.iter_rows(min_row=2, min_col=first_flag, max_col=last_flag):
        for c in row:
            if c.value == 1:
                c.fill = hit

    corr_out = corr.copy()
    corr_out["Phi"] = corr_out["Phi"].round(4)
    corr_out[f"Significant (q < {alpha})"] = np.where(corr_out["q (BH-FDR)"] < alpha, "Yes", "")
    corr_out = corr_out.sort_values(["q (BH-FDR)", "Phi"], ascending=[True, False])
    ws2 = write_sheet(wb, CORR_SHEET, corr_out, {"Category A": 26, "Category B": 26})
    sig_col = corr_out.columns.get_loc(f"Significant (q < {alpha})") + 1
    pos, neg = PatternFill("solid", fgColor="F9D2D2"), PatternFill("solid", fgColor="CDE2FB")
    phi_col = corr_out.columns.get_loc("Phi") + 1
    for row in ws2.iter_rows(min_row=2):
        if row[sig_col - 1].value == "Yes":
            fill = pos if (row[phi_col - 1].value or 0) > 0 else neg
            for c in row:
                c.fill = fill
    wb.save(dst)


# --------------------------------------------------------------------------- explorer page
EXPLORER_ORDER = ["Theme", "Model", "Vehicle", "Country", "Context"]  # panel order in the page


def _first_author(authors):
    parts = [p.strip() for p in str(authors).split(";") if p.strip()] if pd.notna(authors) else []
    if not parts:
        return ""
    return parts[0] + (" et al." if len(parts) > 2 else f" & {parts[1]}" if len(parts) == 2 else "")


def build_explorer(results, ind, cats_by_dim, model_flag, template, out):
    """Write the interactive EV Literature Explorer page: the template with this run's
    classifications embedded as JSON (no extra classification pass)."""
    dims, start = [], 0
    for d in [d for d in EXPLORER_ORDER if d in cats_by_dim]:
        dims.append({"name": d, "cats": cats_by_dim[d], "start": start})
        start += len(cats_by_dim[d])
    cols = [f"{d['name']}: {c}" for d in dims for c in d["cats"]]
    member = ind[cols].values

    def val(r, col, cast=str, default=""):
        v = r.get(col)
        return cast(v) if pd.notna(v) else default

    articles = []
    for i, (_, r) in enumerate(results.iterrows()):
        articles.append([val(r, "Title"), val(r, "Year", int, None), _first_author(r.get("Authors")),
                         val(r, "Source title"), val(r, "Cited by", int, 0), val(r, "DOI"),
                         val(r, "Document Type"), int(model_flag.iloc[i]),
                         [int(k) for k in np.flatnonzero(member[i])]])
    data = {"fields": ["title", "year", "author", "source", "cites", "doi", "type", "model", "cats"],
            "dims": dims, "articles": articles, "n": len(articles)}
    payload = json.dumps(data, ensure_ascii=False, separators=(",", ":")).replace("</", "<\\/")
    html = template.read_text(encoding="utf-8").replace("/*__DATA__*/", payload)
    out.write_text(html, encoding="utf-8")
    print(f"Explorer written to {out.name} ({out.stat().st_size / 1e6:.1f} MB)")


# --------------------------------------------------------------------------- main
def main():
    here = Path(__file__).resolve().parent
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--input", default=here / "Literature review.xlsx", type=Path)
    ap.add_argument("--output", type=Path, help="Workbook to write (default: overwrite --input)")
    ap.add_argument("--charts", default=here / "classification_outputs", type=Path)
    ap.add_argument("--alpha", type=float, default=0.05, help="FDR significance level")
    ap.add_argument("--top", type=int, default=30, help="Pairs shown in the top-correlations chart")
    ap.add_argument("--explorer", default=here / "EV Literature Explorer.html", type=Path,
                    help="Interactive explorer page to write")
    ap.add_argument("--explorer-template", default=here / "explorer_template.html", type=Path)
    ap.add_argument("--no-explorer", action="store_true", help="Skip writing the explorer page")
    args = ap.parse_args()
    output = args.output or args.input
    args.charts.mkdir(parents=True, exist_ok=True)

    # Work from a temporary copy so the source can be read even while open in Excel.
    tmp = Path(tempfile.mkdtemp()) / "source.xlsx"
    shutil.copy2(args.input, tmp)

    results = pd.read_excel(tmp, sheet_name="Results")
    texts = get_texts(results)
    print(f"{len(results):,} articles; {int((texts.str.strip() == '').sum())} with empty text")

    cats_by_dim, flags, readable = {}, [], {}
    for dim in DIMENSIONS:
        cats = load_categories(tmp, dim)
        cats_by_dim[dim] = list(cats)
        print(f"{dim}: {len(cats)} categories, {sum(map(len, cats.values()))} terms")
        ind, matched = classify(texts, cats)
        ind.columns = [f"{dim}: {c}" for c in ind.columns]
        flags.append(ind)
        readable[f"{dim} categories"] = ind.apply(
            lambda r: "; ".join(short(c) for c, v in r.items() if v), axis=1)
        readable[f"{dim} count"] = ind.sum(axis=1)
        readable[f"{dim} matched terms"] = matched
    ind = pd.concat(flags, axis=1)
    dim_of = {c: c.split(":", 1)[0] for c in ind.columns}

    ids = results[[c for c in ID_COLUMNS if c in results]].copy()
    ids.insert(0, "Row", results.index + 2)  # Excel row in Results sheet
    # Flag any word starting with "model" (model, models, modelling, modeling, modelled ...)
    readable[MODEL_FLAG] = texts.str.contains(r"\bmodel\w*", case=False, regex=True).astype(int)
    print(f"{int(readable[MODEL_FLAG].sum()):,} articles mention model* in title, abstract or keywords")
    readable_df = pd.DataFrame(readable)
    order = [f"{d} categories" for d in DIMENSIONS] + [f"{d} count" for d in DIMENSIONS] + [MODEL_FLAG]
    classification = pd.concat(
        [ids, readable_df[order], ind, readable_df[[f"{d} matched terms" for d in DIMENSIONS]]], axis=1)

    corr = pairwise_correlations(ind, dim_of)
    n_sig = int((corr["q (BH-FDR)"] < args.alpha).sum())
    print(f"{len(corr)} category pairs tested; {n_sig} significant at BH-FDR q < {args.alpha}")

    n = len(ind)
    chart_summary(ind, cats_by_dim, n, args.charts / "1_category_counts.png")
    if "Year" in results:
        chart_by_year(ind, results["Year"], cats_by_dim, args.charts / "2_categories_by_year.png")
    chart_crosstabs(ind, cats_by_dim, args.charts / "3_cross_dimension_counts.png")
    chart_correlation_matrix(ind, corr, cats_by_dim, args.alpha, args.charts / "4_correlation_matrix.png")
    chart_top_pairs(corr, args.alpha, args.top, args.charts / "5_top_cross_dimension_correlations.png")
    chart_top_pairs(corr, args.alpha, args.top, args.charts / "6_top_all_correlations.png", cross_only=False)
    classification.to_csv(args.charts / "classification.csv", index=False, encoding="utf-8-sig")
    corr.to_csv(args.charts / "correlations.csv", index=False, encoding="utf-8-sig")
    print(f"Charts and CSVs written to {args.charts}")

    if not args.no_explorer:
        if args.explorer_template.exists():
            try:
                build_explorer(results, ind, cats_by_dim, readable[MODEL_FLAG], args.explorer_template, args.explorer)
            except PermissionError:
                print(f"'{args.explorer.name}' is open in another program - explorer not updated")
        else:
            print(f"Explorer template '{args.explorer_template.name}' not found - explorer not written")

    if output.resolve() == args.input.resolve():
        backup = args.input.with_name(args.input.stem + " (backup).xlsx")
        shutil.copy2(tmp, backup)
        print(f"Backup of original workbook saved as '{backup.name}'")
    try:
        save_workbook(tmp, output, classification, corr, args.alpha)
    except PermissionError:
        alt = output.with_name(output.stem + " (classified).xlsx")
        print(f"'{output.name}' is locked (open in Excel?) - saving to '{alt.name}' instead")
        save_workbook(tmp, alt, classification, corr, args.alpha)
    print("Done.")


if __name__ == "__main__":
    sys.exit(main())
